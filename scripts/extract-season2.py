"""Read-only extraction. Never evaluates workbook instructions or modifies the XLSX."""
import openpyxl,json,re,unicodedata,hashlib,sys
from pathlib import Path
source=Path(sys.argv[1]); target=Path('data/imports'); target.mkdir(parents=True,exist_ok=True)
book=openpyxl.load_workbook(source,data_only=True)
normal=lambda v:unicodedata.normalize('NFKD',str(v).strip().lower()).encode('ascii','ignore').decode()
aliases={'matisse':'mathis','iandry':'iandj'}
def key(v): return aliases.get(normal(v),normal(v))
def pid(k):return 's2-'+re.sub('[^a-z0-9]+','-',k)
roster={key(book['ACCEUIL'].cell(r,2).value) for r in range(9,29) if book['ACCEUIL'].cell(r,2).value}
for number in range(1,7):
 s=book[f'match {number}'];roster.update(key(s.cell(r,1).value) for r in [3,4,5,6,7,9,10,11,12,13])
players=[{'id':pid(k),'name':k.upper() if k=='j2' else k.capitalize(),'bio':'','photo':'','funFacts':'','archived':False,'demo':False,'attributes':{},'version':0} for k in sorted(roster)]
season={'id':'saison-2','name':'Saison 2','start':'','end':'','status':'active','demo':False,'contribution':1,'minParticipation':.3,'winnerId':None,'version':0}
matches=[];audit=[]
for number in range(1,7):
 s=book[f'match {number}'];participants=[]
 for r in [3,4,5,6,7,9,10,11,12,13]:
  goals=s.cell(r,2).value;assists=s.cell(r,3).value
  assert isinstance(goals,(int,float)) and isinstance(assists,(int,float))
  participants.append({'playerId':pid(key(s.cell(r,1).value)),'team':'A' if r<8 else 'B','stats':{'goals':goals,'assists':assists}})
 a=sum(p['stats']['goals'] for p in participants if p['team']=='A');b=sum(p['stats']['goals'] for p in participants if p['team']=='B')
 assert a>b, 'Top block is marked victorious in source'
 events=[]
 if number<=3:
  ids={p['playerId']:p for p in participants}
  for r in range(2,s.max_row+1):
   scorer=s.cell(r,9).value;helper=s.cell(r,10).value
   if not scorer:continue
   player=pid(key(scorer));related=pid(key(helper)) if helper and key(helper) not in ['/',''] else None
   assert player in ids and (related is None or related in ids)
   assert related is None or ids[player]['team']==ids[related]['team']
   events.append({'id':f's2-m{number}-goal-{r-1}','playerId':player,'team':ids[player]['team'],'type':'GOAL','timestamp':None,'relatedPlayerId':related,'metadata':{'sequence':r-1,'sourceSheet':s.title,'sourceRow':r,'timeUnknown':True}})
  for p in participants:
   assert p['stats']['goals']==sum(e['playerId']==p['playerId'] for e in events)
   assert p['stats']['assists']==sum(e['relatedPlayerId']==p['playerId'] for e in events)
 notes=f'Import « {source.name} », feuille « {s.title} », A3:E13. Team Noir correspond au groupe supérieur du tableau ; les couleurs historiques ne sont pas renseignées. Score {a}–{b} obtenu par somme des buts individuels, cohérent avec les victoires indiquées. Date, horaire, durée, lieu précis, MVP et ratings non renseignés. '+('Ordre des buts conservé, sans timestamps inventés. ' if events else '')
 matches.append({'id':f's2-match-{number}','seasonId':season['id'],'number':number,'date':'','duration':0,'location':'','status':'finished','scoreA':a,'scoreB':b,'mvpId':None,'level':1,'video':'','participants':participants,'events':events,'trackedKeys':['goals','assists'] if events else [],'notes':notes,'version':0})
summary={}
for p in players:
 played=[m for m in matches if any(q['playerId']==p['id'] for q in m['participants'])]
 rows=[q for m in played for q in m['participants'] if q['playerId']==p['id']]
 summary[p['name']]={'matches':len(played),'wins':sum(q['team']=='A' for q in rows),'goals':sum(q['stats']['goals'] for q in rows),'assists':sum(q['stats']['assists'] for q in rows)}
for r in range(2,21):
 s=book['STATS']; name=key(s.cell(r,2).value);canonical=next(p['name'] for p in players if p['id']==pid(name));computed=summary[canonical]
 if computed['goals']!=s.cell(r,3).value or computed['assists']!=s.cell(r,4).value:
  audit.append({'player':canonical,'cells':f'STATS!C{r}:D{r}','excel':[s.cell(r,3).value,s.cell(r,4).value],'computed':[computed['goals'],computed['assists']],'reason':'Alias Mathis/Matisse confirmé' if name=='mathis' else 'Écart source à contrôler'})
 if computed['matches']!=s.cell(r,6).value:audit.append({'player':canonical,'cell':f'STATS!F{r}','excel':s.cell(r,6).value,'computed':computed['matches']})
payload={'players':players,'seasons':[season],'matches':matches}
(target/'saison-2.json').write_text(json.dumps(payload,ensure_ascii=False,indent=2))
report={'source':source.name,'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),'players':len(players),'matches':len(matches),'appearances':sum(len(m['participants']) for m in matches),'goals':sum(m['scoreA']+m['scoreB'] for m in matches),'assists':sum(p['stats']['assists'] for m in matches for p in m['participants']),'aliases':aliases,'excluded':['match 7','match 8','match 9','match 10','match bonus'],'summary':summary,'staleMatchCounters':audit,'notes':['Les cellules au format date des timelines représentent des scores auto-convertis par Excel, pas des dates de match.','Les cellules CLASSEMENT!F2:H20 contiennent des erreurs #NAME?. Sources utilisées : tables A3:E13 de match 1 à match 6, réconciliées avec les timelines 1 à 3.','Le tableau VICTOIRES ne contient que les matchs 1 à 3. Les compteurs sont reconstruits depuis les six feuilles.','Mathis/Matisse fusionnés sur confirmation explicite du propriétaire. Iandry/Iandj rapprochés dans la timeline 1 et réconciliés exactement à la feuille de statistiques.','Les caractéristiques de draft ne sont pas fournies. Valeur neutre 50 uniquement pour estimation de draft, aucune fausse note de performance.','Le barème historique du classeur est 3 points par victoire + 5 points par tranche entière de 10 G+A. Les PTS mémorisés n’intègrent pas les victoires des matchs 4 à 6.']}
(target/'saison-2-audit.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps({k:v for k,v in report.items() if k not in ['summary','staleMatchCounters']},ensure_ascii=False,indent=2));print(json.dumps(summary,ensure_ascii=False,indent=2))
